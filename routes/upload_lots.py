"""
routes/upload_lots.py
─────────────────────
POST /api/upload/lots-preview   → parse Excel, return rows for user to review
POST /api/upload/lots-confirm   → actually insert confirmed rows into DB
GET  /api/upload/lots-template  → download blank Excel template
"""
import io
import os
from datetime import datetime, date
from decimal import Decimal
from flask import Blueprint, request, jsonify, send_file
from flask_login import current_user
from db import get_db, ph, is_pg
from routes.auth import approved_required

bp = Blueprint("upload", __name__)

# Exchange → default currency
EXCUR = {
    "NSE": "KES", "NYSE": "USD", "NASDAQ": "USD", "LSE": "GBP",
    "JSE": "ZAR", "EURONEXT": "EUR", "HKEX": "HKD",
    "CRYPTO": "USD", "DSE": "TZS", "USE": "UGX",
    "GSE": "GHS", "BRVM": "XOF",
}

VALID_EXCHANGES = set(EXCUR.keys()) | {"Other"}
VALID_CURRENCIES = {"KES","USD","GBP","EUR","ZAR","TZS","UGX","GHS","HKD","XOF","Other"}

def uid(): return current_user.id
def p():   return ph()


# ── Template download ─────────────────────────────────────────────────────────

@bp.route("/api/upload/lots-template")
@approved_required
def download_template():
    """Serve the pre-built Excel template."""
    template_path = os.path.join(
        os.path.dirname(os.path.dirname(__file__)),
        "static", "lots-import-template.xlsx"
    )
    if not os.path.exists(template_path):
        return jsonify({"error": "Template file not found"}), 404
    return send_file(
        template_path,
        as_attachment=True,
        download_name="lots-import-template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ── Shared parsing logic ──────────────────────────────────────────────────────

def _parse_excel(file_bytes):
    """
    Parse an uploaded Excel file into a list of row dicts.
    Returns: (rows, errors)
      rows   — list of validated dicts ready for insertion
      errors — list of human-readable error strings
    """
    try:
        import openpyxl
    except ImportError:
        return [], ["openpyxl is not installed — contact the admin"]

    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        return [], [f"Could not open file: {e}"]

    # Find the sheet — prefer "Stock Lots", otherwise first sheet
    ws = wb["Stock Lots"] if "Stock Lots" in wb.sheetnames else wb.active

    # Detect header row — look for a row containing "ticker" (case-insensitive)
    header_row_idx = None
    col_map        = {}   # column_name → column_index (0-based)

    for row_idx, row in enumerate(ws.iter_rows(max_row=10, values_only=True)):
        cells = [str(c).lower().strip().replace(" *", "").replace("*", "")
                 if c else "" for c in row]
        if "ticker" in cells:
            header_row_idx = row_idx
            col_map = {name: idx for idx, name in enumerate(cells) if name}
            break

    if header_row_idx is None:
        return [], ["Could not find header row. Make sure the sheet has a 'Ticker' column."]

    required = ["ticker", "exchange", "shares", "purchase price", "date"]
    # Normalise: "purchase price" maps to "purchase_price" key
    alias = {
        "purchase price":  "purchase_price",
        "purchase_price":  "purchase_price",
        "ticker":          "ticker",
        "exchange":        "exchange",
        "shares":          "shares",
        "date":            "date",
        "currency":        "currency",
        "broker":          "broker",
        "note":            "note",
    }

    col_map = {alias.get(k, k): v for k, v in col_map.items() if k in alias}
    missing_cols = [r for r in required if alias.get(r, r) not in col_map]
    if missing_cols:
        return [], [f"Missing required columns: {', '.join(missing_cols)}"]

    rows   = []
    errors = []

    for row_idx, row_vals in enumerate(ws.iter_rows(
            min_row=header_row_idx + 2, values_only=True)):

        # Skip blank rows
        if all(c is None or str(c).strip() == "" for c in row_vals):
            continue

        line = row_idx + header_row_idx + 2  # 1-based line number

        def get(col):
            idx = col_map.get(col)
            if idx is None or idx >= len(row_vals):
                return None
            v = row_vals[idx]
            if v is None:
                return None
            if isinstance(v, (date, datetime)):
                return v.strftime("%Y-%m-%d")
            return str(v).strip()

        # ── Parse each field ────────────────────────────────────────────
        ticker   = (get("ticker") or "").upper().strip()
        exchange = (get("exchange") or "").upper().strip()
        currency = (get("currency") or "").upper().strip()
        broker   = get("broker") or ""
        note     = get("note") or ""

        # Shares
        try:
            shares = float(get("shares") or 0)
            if shares <= 0:
                raise ValueError
        except (ValueError, TypeError):
            errors.append(f"Row {line}: invalid shares '{get('shares')}'")
            continue

        # Purchase price
        try:
            price = float(get("purchase_price") or 0)
            if price <= 0:
                raise ValueError
        except (ValueError, TypeError):
            errors.append(f"Row {line}: invalid purchase price '{get('purchase_price')}'")
            continue

        # Date
        raw_date = get("date") or ""
        parsed_date = None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y",
                    "%Y/%m/%d", "%d %b %Y", "%d %B %Y"):
            try:
                parsed_date = datetime.strptime(raw_date, fmt).strftime("%Y-%m-%d")
                break
            except (ValueError, TypeError):
                continue
        if not parsed_date:
            errors.append(f"Row {line}: invalid date '{raw_date}' — use YYYY-MM-DD")
            continue

        # Required field checks
        if not ticker:
            errors.append(f"Row {line}: ticker is required")
            continue
        if not exchange:
            errors.append(f"Row {line}: exchange is required")
            continue
        if exchange not in VALID_EXCHANGES:
            errors.append(
                f"Row {line}: unknown exchange '{exchange}' "
                f"— use {', '.join(sorted(VALID_EXCHANGES))}")
            continue

        # Auto-fill currency from exchange if blank
        if not currency:
            currency = EXCUR.get(exchange, "KES")
        if currency not in VALID_CURRENCIES:
            currency = EXCUR.get(exchange, "KES")

        rows.append({
            "ticker":         ticker,
            "exchange":       exchange,
            "shares":         shares,
            "purchase_price": price,
            "currency":       currency,
            "date":           parsed_date,
            "broker":         broker,
            "note":           note,
            "lot_value":      round(shares * price, 2),
        })

    return rows, errors


# ── Preview endpoint — parse and return without saving ────────────────────────

@bp.route("/api/upload/lots-preview", methods=["POST"])
@approved_required
def preview_upload():
    """
    Parse the uploaded Excel and return the rows for user review.
    Does NOT save anything to the database.
    """
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    f = request.files["file"]
    if not f.filename:
        return jsonify({"error": "No file selected"}), 400
    if not f.filename.lower().endswith((".xlsx", ".xls")):
        return jsonify({"error": "Please upload an Excel file (.xlsx or .xls)"}), 400

    file_bytes     = f.read()
    rows, errors   = _parse_excel(file_bytes)

    return jsonify({
        "ok":      True,
        "rows":    rows,
        "errors":  errors,
        "count":   len(rows),
        "message": (
            f"Found {len(rows)} lot(s) ready to import"
            + (f" — {len(errors)} row(s) skipped due to errors." if errors else ".")
        ),
    })


# ── Confirm endpoint — actually insert into database ─────────────────────────

@bp.route("/api/upload/lots-confirm", methods=["POST"])
@approved_required
def confirm_upload():
    """
    Insert the rows that were returned by /preview into the database.
    Expects JSON body: {"rows": [...same format as preview response...]}
    """
    d    = request.json or {}
    rows = d.get("rows", [])

    if not rows:
        return jsonify({"error": "No rows to import"}), 400

    conn    = get_db()
    saved   = 0
    errors  = []

    try:
        for r in rows:
            try:
                shares = float(r["shares"])
                price  = float(r["purchase_price"])
                if shares <= 0 or price <= 0:
                    raise ValueError("shares and price must be positive")

                cur = conn.cursor()
                cur.execute(f"""
                    INSERT INTO stock_lots
                        (user_id, ticker, exchange, shares, original_shares,
                         purchase_price, currency, date, broker, note)
                    VALUES ({p()},{p()},{p()},{p()},{p()},{p()},{p()},{p()},{p()},{p()})
                """, (uid(),
                      r["ticker"].upper().strip(),
                      r["exchange"],
                      shares, shares,
                      price,
                      r.get("currency", EXCUR.get(r.get("exchange","NSE"), "KES")),
                      r["date"],
                      r.get("broker", ""),
                      r.get("note", "")))
                saved += 1

            except Exception as e:
                errors.append(f"{r.get('ticker','?')} {r.get('date','?')}: {e}")

        conn.commit()

    except Exception as e:
        conn.rollback()
        return jsonify({"error": f"Database error: {e}"}), 500
    finally:
        conn.close()

    return jsonify({
        "ok":      True,
        "saved":   saved,
        "errors":  errors,
        "message": (
            f"Successfully imported {saved} lot(s)"
            + (f" — {len(errors)} failed." if errors else ".")
        ),
    })
